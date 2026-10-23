"""Checks for the Siena Loop desk comparison workbook."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from real_estate.siena_comps import (
    ASK,
    LIVING_SF,
    NON_AD_VALOREM,
    PRIOR_SALE,
    TOTAL_MILLS,
    build_workbook,
    homestead_tax,
    non_homestead_tax,
    price_per_sf,
    sold_median_ppsf,
    twin_implied_value,
)


def test_address_and_price_per_foot():
    assert price_per_sf(ASK, LIVING_SF) == ASK / 2623
    assert round(price_per_sf(ASK, LIVING_SF), 2) == 455.59
    assert sold_median_ppsf() < 320
    assert twin_implied_value() < 900_000
    assert ASK - twin_implied_value() > 300_000


def test_tax_reset_is_higher_than_the_seller_bill():
    seller_bill = 13_012.60
    assert non_homestead_tax(ASK) == ASK * TOTAL_MILLS / 1000 + NON_AD_VALOREM
    assert non_homestead_tax(ASK) > seller_bill + 7_000
    assert homestead_tax(ASK) > seller_bill + 6_000
    assert homestead_tax(940_000) > seller_bill


def test_workbook_formulas_and_subject_row(tmp_path: Path):
    path = build_workbook(tmp_path / "lakewood-ranch-homes.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "How to use",
        "Subject",
        "Comps",
        "Valuation",
        "Carrying costs",
        "Upgrades",
        "Risks",
        "Sources",
    ]
    subject = wb["Subject"]
    assert subject["B1"].value == "13994 Siena Loop"
    assert "13394" in subject["C1"].value
    assert subject["B4"].value == PRIOR_SALE or subject["B13"].value == PRIOR_SALE
    assert subject["B4"].value == ASK

    comps = wb["Comps"]
    assert comps["B2"].value == "13994 Siena Loop"
    assert comps["E2"].value == ASK
    assert comps["I2"].value == LIVING_SF
    assert comps["M2"].value == "=IF(I2=\"\",\"\",E2/I2)"
    assert comps["B14"].value.startswith("=AGGREGATE")
    assert comps._charts, "price-per-foot chart is missing"
    assert comps["A2"].value == "Subject"
    sold_prices = []
    for row in range(2, 13):
        assert comps.cell(row, 9).value
        if comps.cell(row, 1).value == "Sold":
            sold_prices.append(comps.cell(row, 5).value / comps.cell(row, 9).value)
    assert len(sold_prices) == 6

    valuation = wb["Valuation"]
    assert valuation["B4"].value == ASK
    assert valuation["B6"].value == "=B4/B5"
    assert valuation["B11"].value == "=B10*B5"
    assert valuation["B14"].value == 940_000
    assert valuation["B16"].value == "=B4/B14-1"

    costs = wb["Carrying costs"]
    assert costs["C4"].value == "=Valuation!B4"
    assert costs["C7"].value == "=C5+C6"
    assert "Save Our Homes" in costs["E5"].value

    risks = wb["Risks"]
    severities = [risks.cell(row, 2).value for row in range(2, 16)]
    assert severities.count("High") >= 4
    assert "Roof and insurance" in [risks.cell(row, 1).value for row in range(2, 16)]


def test_2020_sale_multiple():
    assert round(ASK / PRIOR_SALE, 2) == 2.69
