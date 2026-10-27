"""Contract-date summary for 13994 Siena Loop, as a Word document.

The deposit is due Wednesday, October 28, 2026, so the effective date is
Sunday, October 25, 2026. Dates below follow that effective date.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

OUT = ROOT / "real_estate" / "13994-Siena-Loop-contract-dates.docx"
DESKTOP_NAME = "13994 Siena Loop - contract dates.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x1C, 0x20, 0x26)
MUTED = RGBColor(0x4E, 0x56, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BAND = "F1F5F9"
RULE = "D6DCE2"

DATES = [
    ("Wednesday, October 28, 2026", "The $50,000 deposit is due at Premier Settlement Services."),
    (
        "Friday, October 30, 2026",
        "Loan application is due. The seller also delivers any existing title policy and survey.",
    ),
    (
        "Wednesday, November 4, 2026",
        "The 10-day inspection period ends. The written repair notice, any wood-destroying-organism report, and any permit notice are due. Each repair cap is $16,875.",
    ),
    ("Wednesday, November 18, 2026", "The title commitment is due."),
    (
        "Tuesday, November 24, 2026",
        "The 30-day loan approval period ends. Written notice of approval, of a decision to close anyway, or of termination is due before this day.",
    ),
    (
        "Friday, November 27, 2026",
        "If no loan notice was sent, the seller can cancel through this day and return the deposit.",
    ),
    (
        "Monday, November 30, 2026",
        "The seller's deadline to close out permits that were noticed by November 4.",
    ),
    (
        "Wednesday, December 2, or Thursday, December 3, 2026",
        "Walk-through, on December 3 only before the closing time.",
    ),
    (
        "Thursday, December 3, 2026",
        "Closing. The Closing Disclosure waiting period can move it as far as Thursday, December 10, 2026.",
    ),
]


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_run_font(run, name: str, size: float, color: RGBColor, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _paragraph(doc: Document, text: str, *, size: float = 11, color: RGBColor = INK, bold: bool = False, space_after: float = 8, space_before: float = 0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    _set_run_font(run, "Liberation Sans", size, color, bold)


def _heading(doc: Document, text: str) -> None:
    _paragraph(doc, text, size=14, color=NAVY, bold=True, space_before=14, space_after=6)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(0)
        if p.runs:
            p.runs[0].text = item
            _set_run_font(p.runs[0], "Liberation Sans", 11, INK)
        else:
            run = p.add_run(item)
            _set_run_font(run, "Liberation Sans", 11, INK)


def _dates_table(doc: Document) -> None:
    table = doc.add_table(rows=1 + len(DATES), cols=2)
    table.autofit = False
    table.allow_autofit = False
    widths = (Inches(2.7), Inches(4.3))
    headers = ("Date", "What is due")
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.width = widths[index]
        _shade(cell, "1F4E79")
        run = cell.paragraphs[0].add_run(label)
        _set_run_font(run, "Liberation Sans", 10, WHITE, bold=True)
    _repeat_header(table.rows[0])
    _keep_row_together(table.rows[0])
    for row_index, (when, what) in enumerate(DATES, start=1):
        values = (when, what)
        for col, value in enumerate(values):
            cell = table.rows[row_index].cells[col]
            cell.text = ""
            cell.width = widths[col]
            if row_index % 2 == 0:
                _shade(cell, BAND)
            run = cell.paragraphs[0].add_run(value)
            _set_run_font(run, "Liberation Sans", 10, INK, bold=(col == 0))
        _keep_row_together(table.rows[row_index])
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), RULE)
        borders.append(element)
    tbl_pr.append(borders)


def build_docx(path: Path | None = None) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    _paragraph(doc, "13994 Siena Loop", size=22, color=NAVY, bold=True, space_after=2)
    _paragraph(doc, "Contract dates", size=16, color=NAVY, space_after=2)
    _paragraph(
        doc,
        "Lakewood Ranch, Florida 34202  ·  Effective date Sunday, October 25, 2026",
        size=10,
        color=MUTED,
        space_after=10,
    )
    _paragraph(
        doc,
        "The deposit is due Wednesday, October 28, 2026. That fixes the effective date as Sunday, October 25, 2026. The inspection period ends Wednesday, November 4, 2026.",
    )

    _heading(doc, "The contract")
    _paragraph(
        doc,
        "This is the Florida Realtors and Florida Bar Residential Contract For Sale and Purchase, revision 2/26. The As Is rider is unchecked. Paragraph 12 is a repair obligation.",
    )
    _paragraph(
        doc,
        "The purchase price is $1,125,000. The typed $1,000,000 is struck through, and $1,125,000 is written above it and initialed by Antony White, Timothy P. Moriarty, and Shannon L. Moriarty. The initial deposit is $50,000. The typed $25,000 is struck through, and $50,000 is written above it and initialed the same way. It is due within 3 days after the effective date at Premier Settlement Services, 1605 Main Street #1112, Sarasota, FL 34236, (941) 946-8822.",
    )
    _paragraph(
        doc,
        "The loan is 80 percent, which is $900,000, conventional, fixed rate. The term blank means 30 years, at a rate not above the prevailing rate. The buyer pays the owner's title policy. There is no home warranty. Possession is at closing. Closing is Thursday, December 3, 2026.",
    )

    _heading(doc, "Rider B, homeowners' association")
    _paragraph(
        doc,
        "Rider B does not add a cancellation date. The initials, and Antony White's signature at 9:01 AM Pacific on October 24, 2026, acknowledge that the disclosure was in hand before the contract was executed. The three-day cancellation printed on the first page is the remedy when the disclosure arrives after signing. It is not running. Association approval is marked \"is not required,\" so there is no approval deadline and no right to cancel for lack of approval.",
    )
    _paragraph(doc, "Two charges are listed, both through Lakewood Ranch Town Hall (Christine Wofford, 941-907-0202):")
    _bullets(
        doc,
        [
            "$135 per year to Lakewood Ranch Country Club. That is $11.25 a month, the Town Hall figure already in the funding table.",
            "$525 per quarter for ground maintenance. That is $175 a month, $2,100 a year.",
        ],
    )
    _paragraph(
        doc,
        "Together those two lines are $186.25 a month. The funding table uses $186 for Siena plus $11 for Town Hall, or $197. This rider puts the two charges at $186.25 combined. The funding table was not changed. The estoppel from Town Hall, which the seller pays for, is the confirmation before closing. No special assessment amount is filled in. Installments that come due after closing are the buyer's. The community development district assessment stays on the November tax bill and is not one of these two lines.",
    )

    _heading(doc, "Dates")
    _dates_table(doc)

    _heading(doc, "What the November 4 notice has to cover")
    _paragraph(
        doc,
        "The written notice on or before Wednesday, November 4, 2026 preserves the seller's repair duty. If it is late, that duty is waived. The seller still has to maintain the house in the condition it was in on the effective date.",
    )
    _paragraph(
        doc,
        "Each cap is 1.5 percent of $1,125,000 because the dollar blanks are empty: $16,875 for general repairs, $16,875 for wood-destroying organisms, and $16,875 for permits.",
    )
    _paragraph(
        doc,
        "The items that fit the working-condition list are the inoperative irrigation, the loose or broken window sash, and the torn screens. Torn screens are the seller's job before closing. The same notice should include the rotted outdoor-kitchen base, the grill that will not light, the loose front-bedroom trim, and the garage opening. The seller can dispute the items that are not on the working-condition list.",
    )
    _paragraph(
        doc,
        "Cracked roof tiles, with no evidence of an actual leak, are written out of the seller's repair duty. A roofer's written answer still belongs in this window, because an actual leak changes that sentence. An insurance quote belongs in the same window. The contract can move the closing if hazard or wind insurance cannot be issued.",
    )
    _paragraph(
        doc,
        "A price change to $1,100,000 is a separate written amendment. $25,000 is more than the $16,875 repair cap, so the contract does not require the seller to agree. Walking away without a contract right puts the $50,000 deposit at stake.",
    )
    _paragraph(
        doc,
        "The loan in the funding plan is 6.625 percent with a 7-year fixed period and a required refinance, amortized over 30 years. This contract calls for a fixed-rate loan for 30 years at no more than the prevailing rate. That 7-year feature should be matched to Paragraph 8 before Tuesday, November 24, 2026.",
    )

    destination = path or OUT
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return destination


def copy_to_desktop(docx_path: Path, desktop: Path | None = None) -> Path:
    folder = desktop or Path.home() / "Desktop"
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / DESKTOP_NAME
    dest.write_bytes(docx_path.read_bytes())
    return dest


def main() -> None:
    docx_path = build_docx()
    desktop = copy_to_desktop(docx_path)
    print(docx_path)
    print(desktop)


if __name__ == "__main__":
    main()
